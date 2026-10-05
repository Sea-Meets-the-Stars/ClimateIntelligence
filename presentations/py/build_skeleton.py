"""Build the WMKO 2026 skeleton deck on the Kraw_2024 master/theme.

Usage (ocean14 env):
    python presentations/py/build_skeleton.py presentations/2026_WMKO/Kraw_2024.pptx \
        [-o presentations/2026_WMKO/WMKO_2026_Climate_Intelligence.pptx]

Kraw content slides do not use title placeholders: each is the 'TITLE' layout
with its placeholders deleted, a Roboto 38pt bold centered title text box at
the top, and a Roboto 25pt sub-line text box at the bottom. To match exactly,
this script deep-copies those two text boxes from Kraw slide 3 and reuses them
on every new slide. Kraw slide 1 (title 'Sea meets the stars' + sea|sky hero
image) is kept as the title slide; every other Kraw slide is dropped, which
also drops its media.

Slide order, titles and speaker notes come from wmko_skeleton.SLIDES.
"""
import argparse
import copy
import math
import os
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wmko_skeleton import SLIDES, TITLE, DIVIDER  # noqa: E402

DEFAULT_OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "2026_WMKO",
                           "WMKO_2026_Climate_Intelligence.pptx")

CONTENT_LAYOUT = "TITLE"          # what Kraw's content slides use
TITLE_BOX_TEXT = "Deciphering our past"                          # Kraw slide 3 title
SUBLINE_BOX_TEXT = "How did we get here?  What are our origins?"  # Kraw slide 3 sub-line
SLIDE_NUMBER_IDX = 12
TITLE_PT = 38                     # Kraw title size
# Titles in a "group" share one size, the largest that keeps the group's longest
# title on one line (B8b). Measured with Roboto Bold (the Kraw title font,
# OFL-licensed variable font cached in presentations/data/fonts/).
ROBOTO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "fonts", "Roboto.ttf")
TITLE_TEXT_W_IN = 9.3             # 9.75" box minus 0.2" insets, minus a little slack
DIVIDER_TOP_IN = 2.0              # divider title sits mid-slide

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def find_shape(slide, text):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == text:
            return sh
    raise ValueError(f"shape with text {text!r} not found on template slide")


def title_width_in(text, pt):
    """Created by JXP and Claude. Rendered width (inches) of `text` in Roboto
    Bold at `pt` points."""
    font = ImageFont.truetype(ROBOTO, 400)
    font.set_variation_by_name("Bold")
    return font.getlength(text) / 400 * pt / 72


def group_title_sizes(slides):
    """Created by JXP and Claude. Map group name -> title size (whole points):
    the largest size at which every grouped title fits on one line."""
    sizes = {}
    for spec in slides:
        g = spec.get("group")
        if g and spec["title"]:
            w38 = title_width_in(spec["title"], TITLE_PT)
            pt = min(TITLE_PT, math.floor(TITLE_PT * TITLE_TEXT_W_IN / w38))
            sizes[g] = min(sizes.get(g, TITLE_PT), pt)
    # One size for all numbered groups, so titles do not drift between sections.
    common = min(sizes.values(), default=TITLE_PT)
    return {g: common for g in sizes}


def set_size(sp_element, pt):
    """Created by JXP and Claude. Set the font size of every run in a text box."""
    for tag in ("rPr", "endParaRPr"):
        for el in sp_element.iter(f"{A}{tag}"):
            el.set("sz", str(int(pt * 100)))


def set_text(sp_element, text):
    """Replace the text of a text-box element, keeping the first run's formatting."""
    txBody = sp_element.find(f".//{{http://schemas.openxmlformats.org/presentationml/2006/main}}txBody")
    paras = txBody.findall(f"{A}p")
    for p in paras[1:]:
        txBody.remove(p)
    runs = paras[0].findall(f"{A}r")
    for r in runs[1:]:
        paras[0].remove(r)
    for br in paras[0].findall(f"{A}br"):
        paras[0].remove(br)
    runs[0].find(f"{A}t").text = text


def next_shape_id(slide):
    ids = [int(e.get("id")) for e in slide.shapes._spTree.iter() if e.tag.endswith("}cNvPr")]
    return max(ids + [1]) + 1


def add_box(slide, template_el, text, name, pt=None):
    el = copy.deepcopy(template_el)
    cnv = el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
    cnv.set("id", str(next_shape_id(slide)))
    cnv.set("name", name)
    set_text(el, text)
    if pt:
        set_size(el, pt)
    slide.shapes._spTree.append(el)
    return slide.shapes[-1]


def drop_slides(prs, keep):
    sld_ids = prs.slides._sldIdLst
    for i, sld_id in enumerate(list(sld_ids)):
        if i not in keep:
            prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)


def layout_by_name(prs, name):
    for lay in prs.slide_layouts:
        if lay.name == name:
            return lay
    raise ValueError(f"layout {name!r} not in template")


def build(template):
    """Created by JXP and Claude. Build the skeleton deck in memory.
    Input: path to Kraw_2024.pptx. Output: python-pptx Presentation."""
    prs = Presentation(template)
    kraw3 = prs.slides[2]
    title_el = copy.deepcopy(find_shape(kraw3, TITLE_BOX_TEXT)._element)
    subline_el = copy.deepcopy(find_shape(kraw3, SUBLINE_BOX_TEXT)._element)

    drop_slides(prs, keep={0})
    content_layout = layout_by_name(prs, CONTENT_LAYOUT)

    assert SLIDES[0]["kind"] == TITLE, "first skeleton entry must be the title slide"
    title_slide = prs.slides[0]
    for ph in title_slide.placeholders:
        idx = ph.placeholder_format.idx
        if idx == 0:
            set_text(ph._element, SLIDES[0]["title"])
        elif idx == 3 and SLIDES[0]["subtitle"]:
            set_text(ph._element, SLIDES[0]["subtitle"])
    title_slide.notes_slide.notes_text_frame.text = f"[{SLIDES[0]['section']}]\n{SLIDES[0]['notes']}"

    sizes = group_title_sizes(SLIDES)
    print("group title sizes (pt):", sizes)
    for spec in SLIDES[1:]:
        slide = prs.slides.add_slide(content_layout)
        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx != SLIDE_NUMBER_IDX:
                ph._element.getparent().remove(ph._element)
        if spec["title"]:
            box = add_box(slide, title_el, spec["title"], "Title", sizes.get(spec.get("group")))
            if spec["kind"] == DIVIDER:
                box.top = Inches(DIVIDER_TOP_IN)
        if spec["subtitle"]:
            add_box(slide, subline_el, spec["subtitle"], "Sub-line")
        slide.notes_slide.notes_text_frame.text = f"[{spec['section']}]\n{spec['notes']}"
    return prs


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("template")
    p.add_argument("-o", "--output", default=DEFAULT_OUT)
    args = p.parse_args()

    prs = build(args.template)
    out = os.path.abspath(args.output)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    print(f"wrote {out}: {len(prs.slides)} slides, {os.path.getsize(out)/1e6:.2f} MB")


if __name__ == "__main__":
    main()
