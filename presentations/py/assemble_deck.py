"""Created by JXP and Claude.

Assemble the WMKO 2026 deck: rebuild the skeleton on the Kraw master
(build_skeleton.build), then place each slide's figure, instrument photo and
~16pt source line (Q10), and append the full figure source to the speaker
notes (from figs/sources.json).

B8 = slide 4 + the climate half (C1-C10); B9 = intro (slides 3, 5) + the AI
half (A1-A10). Images are looked up in figs/, then figs/images/, then relative
to the repo root (sacbee.png, docs/CI_graphic.png).

Layout (slide 10" x 5.625"; nothing within 0.4" of the side edges):
  - content area starts below the title (1 line ~0.95", 2 lines ~1.5");
  - source text sits above the master's UCSC logo (y >= 5.24") and left of
    the slide number (x >= 9.27");
  - a figure with an instrument photo gets the left ~6.4"; the photo the
    right ~2.6"; both scaled to fit, aspect preserved, vertically centered.

Usage:
    conda run -n ocean14 python presentations/py/assemble_deck.py [-o OUT.pptx]
"""
import argparse
import json
import os
import tempfile
from pathlib import Path

from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from build_skeleton import build, title_width_in, TITLE_PT, TITLE_TEXT_W_IN
from slide_style import FIGS, PRES, ROOT

TEMPLATE = PRES / "2026_WMKO" / "Kraw_2024.pptx"
OUT = PRES / "2026_WMKO" / "WMKO_2026_Climate_Intelligence.pptx"
IMG = FIGS / "images"

X0, X1 = 0.4, 9.6            # usable width (inches)
SRC_X, SRC_W = 0.4, 8.75     # source line box (ends before the slide number)
SRC_TOP_1LINE = 4.92         # top of a one-line source; content ends above it
SRC_LINE_H = 0.27            # 16 pt line height (inches)
SRC_CHARS_PER_LINE = 80      # ~16 pt Roboto across SRC_W
PHOTO_W = 2.6
GAP = 0.2
NOTE_PT = 18                 # slide-level explanatory note (B8b), above the source line
NOTE_H = 0.62                # two lines at 18 pt

# Slide title -> short explanatory note placed under the figure (numbers from heat_scale.py).
NOTES = {
    "3: Heat in vs. heat out": (
        "Too little to feel (~2% of your resting body's heat). "
        "But 0.7 W/m² over all of Earth, nonstop, is ~360 TW: 18× all human energy use."),
    # Q26a: Gates's words carry the bio claim (Fortune, 30 Sep 2026).
    "2: AI can already help create a deadly pandemic": (
        "“A.I. has crossed the threshold that its ability to empower a bioterrorist to kill "
        "hundreds of millions — that exists today.” — Bill Gates"),
}

# A7: two photos side by side, each with a caption (Q30, Q40, Q45).
PAIRS = {
    "7: AI will steeply raise the value of data collection": (
        ("keck_primary.jpg", "Keck: 36 hexagonal segments, a 10 m mirror"),
        ("hires.jpg", "HIRES (PI S. Vogt): built 1988–93, first light 16 Jul 1993"),
        "Photos: z2amiller / Wikimedia Commons (CC BY-SA 2.0); courtesy W. M. Keck Observatory · "
        "Vogt et al. 1994, Proc. SPIE 2198"),
}
# Slides whose image already carries its own title: drop the slide title so the image fills the slide.
HIDE_TITLE = set()

# Slides whose source line sits low and right-aligned (B9b: give the figure more room).
# The source is sized (<= 16 pt, >= 12 pt) to fit on one line, measured in Roboto.
SOURCE_LOW_RIGHT = {"2: AI can already help create a deadly pandemic"}
LOW_SRC_X, LOW_SRC_W, LOW_SRC_TOP = 1.95, 7.2, 5.15   # right of the logo, left of the slide number

# Wider photo/inset column for specific slides (B9b: fill the white space on A4).
PHOTO_W_FOR = {"4: Resistance to data centers is a (valuable) distraction": 3.1}

# A10: a labeled empty box; the author adds the image (Q43).
PLACEHOLDERS = {
    "10: The Eye of Sauron": "Eye of Sauron image — to be added (see To Do)",
}

# Slide title (exact skeleton title) -> (figure, instrument photo/inset or None, slide source line or None)
PLACEMENTS = {
    "The two exponentials": (
        "s4_two_exponentials.png", None,
        "Data: Global Carbon Project via Our World in Data (left); Epoch AI, Notable AI Models (right)"),
    "1: CO₂ has been rising for 50+ years": (
        "c1_keeling.png", "mauna_loa_observatory.jpg",
        "Data: NOAA GML / Scripps, Mauna Loa Observatory · Photo: Brian Vasel/NOAA"),
    "2: The U.S. has dominated greenhouse gas release": (
        "c2_us_cumulative.png", None,
        "Data: Global Carbon Project, Global Carbon Budget 2025, via Our World in Data"),
    "3: Heat in vs. heat out": (
        "c3a_energy_budget.png", "ceres_instrument.jpg",
        "Data: IPCC AR6 WGI Fig. 7.2; CERES satellites (Loeb et al. 2021) · Photo: NASA"),
    "3: Heat in vs. heat out: the CO₂ bite": (
        "c3b_olr_spectrum.png", None,
        "Schematic from Planck's law; shape after satellite spectra (Nimbus-4 IRIS, Hanel et al. 1972)"),
    "4: The ocean is warming: surface": (
        "c4a_sst_global.png", None,
        "Data: NOAA NCEI, NOAAGlobalTemp / ERSSTv5 (ships, buoys, Argo)"),
    "4: The ocean is warming: depth": (
        "c4b_ohc_by_depth.png", "argo_float.jpg",
        "Data: NOAA NCEI ocean heat content (XBT/CTD, then Argo) · Photo: NOAA GOMO"),
    "5: The ocean isn't done warming": (
        "c5_committed_warming.png", None,
        "IPCC AR6 WGI (SPM A.1.2, B.5.4; Ch. 4, 7); Murphy 2021"),
    "6: Sea level is rising": (
        "c6a_gmsl.png", "sentinel6.jpg",
        "Tide gauges: CSIRO (Church & White 2011). Altimetry data are provided by NOAA "
        "Laboratory for Satellite Altimetry · Image: NASA"),
    "6: Sea level is rising: Honolulu": (
        "c6b_honolulu.png", "honolulu_tide_station.jpg",
        "Data and photo: NOAA CO-OPS, Honolulu Harbor tide gauge 1612340"),
    "7: Biodiversity is declining fast": (
        "c7a_lpi_biomass.png", None,
        "Data: WWF/ZSL Living Planet Index 2024 via Our World in Data; Bar-On et al. 2018 (PNAS)"),
    "7: Biodiversity is declining fast: extinctions": (
        "c7b_extinctions.png", None,
        "Data: Ceballos et al. 2015, Science Advances (IUCN Red List)"),
    "8: Planetary boundaries are being exceeded": (
        "c8_planetary_bounds.png", None,
        "Data: Richardson et al. 2023 (Science Advances); Planetary Health Check 2025"),
    "9: Population growth is poised to decline": (
        "c9a_population.png", None,
        "Data: UN World Population Prospects 2024 via Our World in Data; IHME (Vollset et al. 2020)"),
    "9: Fertility has fallen below replacement": (
        "c9b_fertility.png", None,
        "Data: UN World Population Prospects 2024 via Our World in Data"),
    "10: Most economists believe technology will save us": (
        "c10_economist_survey.png", None,
        "Survey: Howard & Sylvan 2021, Institute for Policy Integrity (738 climate economists, worldwide)"),
    # --- B9: intro ---
    "Climate Intelligence": ("s3_ci_concept.png", None, None),
    "Our own sensors are exquisite": (
        "s4_senses.png", None,
        "Vision: Hecht et al. 1942 · hearing: 0–120 dB · geosmin ~5 ng/L (USGS) · touch: Skedung et al. 2013"),
    "Humans suck at exponentials": (
        "s5_covid.png", None, "Data: Johns Hopkins CSSE via Our World in Data (Blog 001, Figure 1)"),
    # --- B9: AI half ---
    "1: AI has surpassed me as a scientist": (
        "presentations/2026_WMKO/sacbee.png", None,
        "J. Xavier Prochaska, The Sacramento Bee, 21 Aug 2026 (photo: Dado Ruvic/Reuters)"),
    "2: AI can already help create a deadly pandemic": (
        "a2_bio_uplift.png", None,
        "Fortune 30 Sep 2026; RAND, OpenAI 2024; Anthropic 2025; Zhang+ 2026 (not wet-lab)"),
    "3: AI can already hack nearly every system on Earth": (
        "a3_hacking.png", None,
        "Scientific American & Axios, Apr 2026 (Mythos, UK AISI); CNBC, Jul 2026 (Hugging Face)"),
    "4: Resistance to data centers is a (valuable) distraction": (
        "a4_dc_electricity.png", "a4_dc_water_inset.png",
        "Data: LBNL 2024 U.S. Data Center Energy Usage Report and 2025 Update (electricity, water)"),
    "5: AI will force the collapse of peer review": (
        "a5_arxiv.png", None, "Data: arXiv monthly submissions; arXiv blog, 1 Oct 2026"),
    "6: AI may end grant competition": ("a6_proposals.png", None, None),
    "8: The skills humans need": ("a8_skills.png", None, None),
    "9: The AI bubble will pop (with a few major winners)": (
        "a9_ai_capex.png", None,
        "Data: 10-K filings via SEC EDGAR (MSFT fiscal years); 2026 = July 2026 guidance"),
}


def title_bottom(title, pt):
    """Created by JXP and Claude. Bottom (inches) of a Kraw title box: 0.02"
    offset + 0.2" insets + line height, with the line count measured in
    Roboto Bold at the title's actual size."""
    lines = 1 if title_width_in(title, pt) <= TITLE_TEXT_W_IN else 2
    return 0.02 + 0.2 + lines * pt * 1.2 / 72 + 0.06


MAX_PX = 2000                 # photos larger than this are downsampled for the deck (B10)
SLIDE_IMG_CACHE = Path(tempfile.gettempdir()) / "wmko2026_slide_images"


def resolve(name):
    """Created by JXP and Claude. Path of an image: figs/, figs/images/, or
    repo-relative. Photos/screenshots with a side > MAX_PX are downsampled to a
    JPEG copy (originals untouched) to keep the deck small."""
    for base in (FIGS, IMG, ROOT):
        path = base / name
        if path.exists():
            break
    else:
        raise FileNotFoundError(name)
    if base == FIGS:
        return path  # our own figures are already slide-sized
    with Image.open(path) as im:
        if max(im.size) <= MAX_PX and path.stat().st_size < 600_000:
            return path
        SLIDE_IMG_CACHE.mkdir(exist_ok=True)
        out = SLIDE_IMG_CACHE / (path.stem + "_slide.jpg")
        small = im.convert("RGB")
        small.thumbnail((MAX_PX, MAX_PX))
        small.save(out, quality=88)
    return out


def fit(path, box_w, box_h):
    """Created by JXP and Claude. Largest (w, h) in inches with the image's
    aspect ratio that fits inside box_w x box_h."""
    with Image.open(path) as im:
        aspect = im.width / im.height
    w = min(box_w, box_h * aspect)
    return w, w / aspect


def text_width_in(text, pt):
    """Created by JXP and Claude. Width (inches) of text in Roboto Regular."""
    from PIL import ImageFont
    font = ImageFont.truetype(str(PRES / "data" / "fonts" / "Roboto.ttf"), 400)
    font.set_variation_by_name("Regular")
    return font.getlength(text) / 400 * pt / 72


def add_text(slide, text, top, height, pt, color, name, center=False, x=SRC_X, w=SRC_W, right=False):
    """Created by JXP and Claude. A plain Roboto text box (default: across the slide)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(top), Inches(w), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    para = tf.paragraphs[0]
    if center:
        para.alignment = PP_ALIGN.CENTER
    elif right:
        para.alignment = PP_ALIGN.RIGHT
    run = para.add_run()
    run.text = text
    run.font.size = Pt(pt)
    run.font.name = "Roboto"
    run.font.color.rgb = color
    tb.name = name


def add_source(slide, text, top, n_lines):
    """Created by JXP and Claude. ~16 pt gray source line in the Kraw font."""
    add_text(slide, text, top, SRC_LINE_H * n_lines + 0.05, 16, RGBColor(0x66, 0x66, 0x66), "Source")


def place(slide, title, title_pt, fig_name, photo_name, source):
    """Created by JXP and Claude. Lay out one slide's figure, photo, note, source."""
    top = 0.4 if title in HIDE_TITLE else title_bottom(title, title_pt)
    low = source and title in SOURCE_LOW_RIGHT
    if low:
        pt = next((p for p in (16, 15, 14, 13, 12) if text_width_in(source, p) <= LOW_SRC_W - 0.05), 12)
        add_text(slide, source, LOW_SRC_TOP, SRC_LINE_H, pt, RGBColor(0x66, 0x66, 0x66), "Source",
                 x=LOW_SRC_X, w=LOW_SRC_W, right=True)
        n_src, src_top = 0, LOW_SRC_TOP + 0.02
    else:
        n_src = -(-len(source) // SRC_CHARS_PER_LINE) if source else 0
        src_top = SRC_TOP_1LINE - SRC_LINE_H * (n_src - 1) if source else SRC_TOP_1LINE + 0.2
    bottom = src_top - 0.08
    note = NOTES.get(title)
    if note:
        bottom -= NOTE_H + 0.05
        add_text(slide, note, bottom + 0.05, NOTE_H, NOTE_PT, RGBColor(0x22, 0x22, 0x22), "Note", center=True)
    avail_h = bottom - top

    photo_w = PHOTO_W_FOR.get(title, PHOTO_W)
    fig_box_w = (X1 - X0) - (photo_w + GAP if photo_name else 0)
    fig_path = resolve(fig_name)
    fw, fh = fit(fig_path, fig_box_w, avail_h)
    fx = X0 + (fig_box_w - fw) / 2
    slide.shapes.add_picture(str(fig_path), Inches(fx), Inches(top + (avail_h - fh) / 2),
                             Inches(fw), Inches(fh)).name = "Figure"
    if photo_name:
        photo_path = resolve(photo_name)
        pw, ph = fit(photo_path, photo_w, avail_h * (1.0 if title in PHOTO_W_FOR else 0.9))
        slide.shapes.add_picture(str(photo_path), Inches(X1 - photo_w + (photo_w - pw) / 2),
                                 Inches(top + (avail_h - ph) / 2), Inches(pw), Inches(ph)).name = "Instrument"
    if source and not low:
        add_source(slide, source, src_top, n_src)


def place_pair(slide, title, title_pt, left, right, source):
    """Created by JXP and Claude. Two captioned photos side by side (A7)."""
    top = title_bottom(title, title_pt)
    n_src = -(-len(source) // SRC_CHARS_PER_LINE)
    src_top = SRC_TOP_1LINE - SRC_LINE_H * (n_src - 1)
    cap_h = 0.4
    avail_h = src_top - 0.1 - cap_h - top
    col_w = (X1 - X0 - GAP) / 2
    for k, (img, caption) in enumerate((left, right)):
        path = resolve(img)
        w, h = fit(path, col_w, avail_h)
        x = X0 + k * (col_w + GAP) + (col_w - w) / 2
        y = top + (avail_h - h) / 2
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h)).name = "Photo"
        tb = slide.shapes.add_textbox(Inches(X0 + k * (col_w + GAP)), Inches(y + h + 0.05),
                                      Inches(col_w), Inches(cap_h))
        tb.text_frame.word_wrap = True
        para = tb.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = caption
        run.font.size = Pt(16)
        run.font.name = "Roboto"
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        tb.name = "Caption"
    add_source(slide, source, src_top, n_src)


def place_placeholder(slide, title, title_pt, label):
    """Created by JXP and Claude. A dashed, labeled empty box for an image the
    author will add (A10)."""
    top = title_bottom(title, title_pt) + 0.1
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(top),
                                 Inches(7.0), Inches(SRC_TOP_1LINE - top - 0.1))
    box.fill.background()
    box.line.color.rgb = RGBColor(0x99, 0x99, 0x99)
    box.line.width = Pt(1.5)
    box.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
    tf = box.text_frame
    tf.text = label
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(18)
    run.font.name = "Roboto"
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    box.name = "Image placeholder"


def main():
    """Created by JXP and Claude. Rebuild the deck and place all content."""
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("-o", "--output", default=str(OUT))
    args = p.parse_args()

    sources = json.loads((FIGS / "sources.json").read_text())
    prs = build(TEMPLATE)
    placed = set()
    for i, slide in enumerate(prs.slides, start=1):
        titles = [sh.text_frame.text for sh in slide.shapes if sh.name == "Title"]
        if not titles:
            continue
        title = titles[0]
        title_shape = next(sh for sh in slide.shapes if sh.name == "Title")
        size = title_shape.text_frame.paragraphs[0].runs[0].font.size
        title_pt = size.pt if size else TITLE_PT
        if title in PAIRS:
            left, right, source = PAIRS[title]
            place_pair(slide, title, title_pt, left, right, source)
            what = f"{left[0]} + {right[0]}"
        elif title in PLACEHOLDERS:
            place_placeholder(slide, title, title_pt, PLACEHOLDERS[title])
            what = "placeholder box"
        elif title in PLACEMENTS:
            fig_name, photo_name, source = PLACEMENTS[title]
            place(slide, title, title_pt, fig_name, photo_name, source)
            if title in HIDE_TITLE:
                title_shape._element.getparent().remove(title_shape._element)
            if fig_name in sources:
                notes = slide.notes_slide.notes_text_frame
                notes.text = notes.text + f"\n\nFull figure source ({fig_name}): {sources[fig_name]['source']}"
            what = fig_name + (f" + {photo_name}" if photo_name else "")
        else:
            continue
        placed.add(title)
        print(f"slide {i:2d}: {title[:45]:45s} <- {what}")
    missing = (set(PLACEMENTS) | set(PAIRS) | set(PLACEHOLDERS)) - placed
    assert not missing, f"titles not found in skeleton: {missing}"

    prs.save(args.output)
    print(f"wrote {args.output}: {len(prs.slides)} slides, {os.path.getsize(args.output) / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
